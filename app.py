"""
Prism AI — agentic chatbot backend (Flask + NVIDIA inference API + Claude API).

Capabilities: streaming chat, native tool-calling agent loop (web search, page
reading, Python execution, workspace file access, shell commands, image
generation, persistent memory), file/PDF analysis, vision input.
"""

import csv
import io
import json
import os
import time
import uuid

import httpx
from dotenv import load_dotenv
from flask import (
    Flask,
    Response,
    jsonify,
    render_template,
    request,
    send_from_directory,
    stream_with_context,
)

from agent.loop import run as run_agent_loop
from agent.memory import add_memory, init_db
from agent.prompt import system_prompt
from agent.timing import (
    bind_request,
    clear_request,
    elapsed_ms,
    end_request,
    record,
    start_request,
)
from agent.tools import WORKSPACE, set_memory_writer, t_run_python, tool_defs
from prism_superagent.api_registry.errors import APIProvidersUnavailable
from prism_superagent.api_registry.integration import (
    classify_chat_routes,
    select_chat_provider,
)
from prism_superagent.api_registry.runtime import APIRegistryRuntime, format_response
from prism_superagent.capabilities.ollama import OllamaModelRegistry
from prism_superagent.capabilities.registry import build_capability_registry
from prism_superagent.capabilities.summarization import LocalDocumentSummarizer
from prism_superagent.engine.attachments import attached_documents, request_text
from prism_superagent.engine.documents import ocr_pdf_bytes
from prism_superagent.engine.errors import FallbackToModel
from prism_superagent.engine.runtime import DeterministicEngine
from prism_superagent.tracing import ExecutionTrace, bind_trace, reset_trace
from providers import anthropic as anthropic_provider
from providers import openai_compatible as nvidia_provider

load_dotenv()

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

VISION_MODEL = "meta/llama-3.2-90b-vision-instruct"

MODELS = [
    {
        "id": "moonshotai/kimi-k3",
        "name": "Kimi K3",
        "tag": "best overall",
        "provider": "nvidia",
    },
    {
        "id": "deepseek-ai/deepseek-v4-pro-0813",
        "name": "DeepSeek V4 Pro",
        "tag": "reasoning",
        "provider": "nvidia",
    },
    {
        "id": "deepseek-ai/deepseek-v4-flash-0731",
        "name": "DeepSeek V4 Flash",
        "tag": "fast",
        "provider": "nvidia",
    },
    {
        "id": "nvidia/nemotron-3-ultra-550b-a55b",
        "name": "Nemotron 3 Ultra 550B",
        "tag": "agent",
        "provider": "nvidia",
    },
    {
        "id": "nvidia/nemotron-3.5-lightning-30b-a3b",
        "name": "Nemotron 3.5 Lightning",
        "tag": "fast",
        "provider": "nvidia",
    },
    {
        "id": "nvidia/nemotron-3-super-120b-a12b",
        "name": "Nemotron 3 Super 120B",
        "tag": "balanced",
        "provider": "nvidia",
    },
    {
        "id": "openai/gpt-oss-20b",
        "name": "GPT-OSS 20B",
        "tag": "open source",
        "provider": "nvidia",
    },
    {
        "id": "mistralai/mistral-large-2-instruct",
        "name": "Mistral Large 2",
        "tag": "multilingual",
        "provider": "nvidia",
    },
    {
        "id": "mistralai/codestral-22b-instruct-v0.1",
        "name": "Codestral 22B",
        "tag": "code",
        "provider": "nvidia",
    },
    {
        "id": VISION_MODEL,
        "name": "Llama 3.2 90B Vision",
        "tag": "images",
        "provider": "nvidia",
    },
    {
        "id": "claude-opus-5",
        "name": "Claude Opus 5",
        "tag": "deepest reasoning",
        "provider": "anthropic",
    },
    {
        "id": "claude-sonnet-5",
        "name": "Claude Sonnet 5",
        "tag": "flagship, balanced",
        "provider": "anthropic",
    },
]

init_db()
set_memory_writer(add_memory)
deterministic_engine = DeterministicEngine(WORKSPACE)
api_registry_runtime = APIRegistryRuntime()
ollama_model_registry = OllamaModelRegistry(
    os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434")
)
local_document_summarizer = LocalDocumentSummarizer(
    os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434")
)
capability_registry = build_capability_registry(
    deterministic_engine,
    api_registry_runtime.registry,
    agent_tools=tool_defs(agent_mode=True),
    cloud_models=MODELS,
)


# ---------------------------------------------------------------- routes
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/health")
def health():
    return jsonify({"status": "ok"})


@app.teardown_request
def _clear_timing_context(_error=None):
    clear_request()


@app.route("/files/<path:relpath>")
def files(relpath):
    return send_from_directory(WORKSPACE, relpath)


@app.route("/api/config")
def config():
    ollama_model_registry.refresh()
    _refresh_capability_registry()
    key = os.getenv("NVIDIA_API_KEY") or ""
    return jsonify(
        {
            "models": MODELS,
            "hasKey": bool(key),
            "keyLooksNvidia": key.startswith("nvapi-"),
        }
    )


@app.route("/api/chat", methods=["POST"])
def chat():
    request_id = uuid.uuid4().hex[:12]
    _, request_started = start_request(request_id)
    data = request.get_json(force=True)
    messages = data.get("messages", [])
    requested_model = data.get("model")
    automatic_routing = requested_model in (None, "", "auto")
    model = MODELS[0]["id"] if automatic_routing else requested_model
    agent_mode = bool(data.get("agent"))
    web_mode = bool(data.get("web"))

    has_image = any(isinstance(m.get("content"), list) for m in messages)
    if has_image:
        model = VISION_MODEL  # Claude models don't take image input here; NVIDIA vision model covers it

    last_user = next((m for m in reversed(messages) if m.get("role") == "user"), {})
    deterministic_graph, api_match = classify_chat_routes(
        last_user.get("content"),
        deterministic_engine,
        api_registry_runtime,
        web_mode=web_mode,
    )
    summary_documents = []
    selected_local_model = None
    user_content = last_user.get("content")
    if (
        automatic_routing
        and not deterministic_graph
        and not api_match
        and not has_image
        and isinstance(user_content, str)
    ):
        summary_request = request_text(user_content)
        if attached_documents(user_content) and _is_document_summary_request(
            summary_request
        ):
            ollama_model_registry.refresh()
            _refresh_capability_registry()
            candidates = tuple(
                item
                for item in ollama_model_registry.confirmed_local_models("completion")
                if capability_registry.get(
                    f"local_model:{item.name}:document_summarization"
                )
            )
            if candidates:
                summary_documents = attached_documents(user_content)
                selected_local_model = candidates[0]
    selected_provider_name = next(
        (m["provider"] for m in MODELS if m["id"] == model), "nvidia"
    )
    fallback_provider = (
        anthropic_provider if selected_provider_name == "anthropic" else nvidia_provider
    )
    provider = (
        "local_model"
        if selected_local_model
        else select_chat_provider(
            deterministic_graph, api_match, selected_provider_name
        )
    )
    trace_capability = (
        deterministic_graph.kind
        if deterministic_graph
        else api_match.capability
        if api_match
        else "semantic_document_summary"
        if selected_local_model
        else "agent_tool_workflow"
        if agent_mode
        else "general_chat"
    )
    trace_request = _trace_request_text(user_content)
    trace = ExecutionTrace(trace_request, agent_mode=agent_mode)
    trace_model = (
        selected_local_model.name
        if selected_local_model
        else model
        if provider in ("nvidia", "anthropic")
        else None
    )
    trace.set_route(provider, trace_capability, trace_model)
    record("request_route", provider=provider, model=model)
    system_text = (
        "" if deterministic_graph or api_match else system_prompt(agent_mode, web_mode)
    )
    convo = [{"role": "system", "content": system_text}] + messages
    tools = None if has_image else tool_defs(agent_mode)  # vision model: plain chat

    def sse(obj):
        if obj.get("done") is True or obj.get("error"):
            obj = {**obj, "execution_trace": trace.snapshot()}
        return "data: " + json.dumps(obj, ensure_ascii=False) + "\n\n"

    def generate_claude():
        yield from run_agent_loop(
            anthropic_provider, model, system_text, convo, tools, sse
        )

    def generate():
        if deterministic_graph:
            started = time.perf_counter()
            yield sse(
                {"tool": "deterministic_engine", "label": deterministic_graph.kind}
            )
            try:
                result = deterministic_engine.execute(deterministic_graph)
            except FallbackToModel:
                record(
                    "deterministic_escalation",
                    task_id=deterministic_graph.task_id,
                    kind=deterministic_graph.kind,
                    provider=selected_provider_name,
                )
                yield sse(
                    {
                        "notice": "The local retrieval route was unavailable; continuing with the selected provider."
                    }
                )
                fallback_system = system_prompt(agent_mode, web_mode)
                convo[0]["content"] = fallback_system
                yield from run_agent_loop(
                    fallback_provider, model, fallback_system, convo, tools, sse
                )
                return
            except Exception as e:
                record(
                    "deterministic_task_error",
                    task_id=deterministic_graph.task_id,
                    kind=deterministic_graph.kind,
                    error_type=type(e).__name__,
                )
                yield sse({"tool_done": "deterministic_engine"})
                yield sse({"error": f"Prism's local task could not be completed: {e}"})
                return
            record(
                "deterministic_task_end",
                task_id=result.task_id,
                kind=result.kind,
                cache_hit=result.cache_hit,
                duration_ms=elapsed_ms(started),
                verified=True,
                result_status=result.result_status,
            )
            yield sse({"tool_done": "deterministic_engine"})
            yield sse({"content": result.answer})
            yield sse({"done": True})
            return
        if api_match:
            started = time.perf_counter()
            yield sse({"tool": "api_registry", "label": api_match.capability})
            try:
                result = api_registry_runtime.execute(api_match)
            except APIProvidersUnavailable as e:
                record(
                    "api_registry_fallback",
                    capability=e.capability,
                    failed_providers=[item[0] for item in e.failures],
                    duration_ms=elapsed_ms(started),
                )
                yield sse({"tool_done": "api_registry"})
                yield sse(
                    {
                        "notice": "No configured API provider completed the request; continuing with the selected provider."
                    }
                )
                fallback_system = system_prompt(agent_mode, web_mode)
                convo[0]["content"] = fallback_system
                yield from run_agent_loop(
                    fallback_provider, model, fallback_system, convo, tools, sse
                )
                return
            except Exception as e:
                record(
                    "api_registry_error",
                    capability=api_match.capability,
                    error_type=type(e).__name__,
                    duration_ms=elapsed_ms(started),
                )
                yield sse({"tool_done": "api_registry"})
                yield sse(
                    {
                        "notice": "The API registry could not complete this request; continuing with the selected provider."
                    }
                )
                fallback_system = system_prompt(agent_mode, web_mode)
                convo[0]["content"] = fallback_system
                yield from run_agent_loop(
                    fallback_provider, model, fallback_system, convo, tools, sse
                )
                return
            record(
                "api_registry_success",
                provider_id=result.provider_id,
                capability=result.capability,
                cache_hit=result.cache_hit,
                latency_ms=elapsed_ms(started),
                http_status=result.status,
                response_sha256=result.sha256,
                result_status=result.result_status,
            )
            yield sse({"tool_done": "api_registry"})
            yield sse({"content": format_response(result)})
            yield sse({"done": True})
            return
        if selected_local_model:
            started = time.perf_counter()
            yield sse({"tool": "local_model", "label": "document summarization"})
            try:
                result = local_document_summarizer.summarize(
                    selected_local_model, summary_documents, summary_request
                )
            except (httpx.HTTPError, ValueError) as e:
                record(
                    "local_summary_fallback",
                    model=selected_local_model.name,
                    error_type=type(e).__name__,
                    duration_ms=elapsed_ms(started),
                )
                yield sse({"tool_done": "local_model"})
                yield sse(
                    {
                        "notice": "Local summarization was unavailable; continuing with the selected provider."
                    }
                )
                yield from run_agent_loop(
                    fallback_provider, model, system_text, convo, tools, sse
                )
                return
            record(
                "local_summary_end",
                model=result.model,
                status=result.status,
                verification_status=result.verification_status,
                duration_ms=elapsed_ms(started),
            )
            yield sse({"tool_done": "local_model"})
            yield sse(
                {
                    "notice": f"Local model summary ({result.status}; {result.verification_status}; semantic accuracy is not independently verified)."
                }
            )
            yield sse({"content": result.answer})
            yield sse({"done": True})
            return
        yield from run_agent_loop(
            nvidia_provider, model, system_text, convo, tools, sse
        )

    gen = (
        generate_claude
        if provider == "anthropic" and not (deterministic_graph or api_match)
        else generate
    )

    def timed_stream():
        trace_token = bind_trace(trace)
        stream_timing_token = bind_request(request_id, request_started)
        try:
            yield from gen()
        finally:
            record("request_end", total_ms=elapsed_ms(request_started))
            end_request(stream_timing_token)
            reset_trace(trace_token)

    return Response(
        stream_with_context(timed_stream()),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


def _is_document_summary_request(text):
    import re

    return bool(
        re.search(
            r"\b(?:summar(?:y|ize|ise|izing|ising|ization)|overview)\b",
            text,
            re.IGNORECASE,
        )
    )


def _trace_request_text(content):
    if isinstance(content, str):
        return request_text(content)
    if isinstance(content, list):
        return " ".join(
            item.get("text", "")
            for item in content
            if isinstance(item, dict) and isinstance(item.get("text"), str)
        )
    return ""


def _refresh_capability_registry():
    global capability_registry
    capability_registry = build_capability_registry(
        deterministic_engine,
        api_registry_runtime.registry,
        agent_tools=tool_defs(agent_mode=True),
        cloud_models=MODELS,
        local_models=ollama_model_registry.models,
    )


@app.route("/api/run_code", methods=["POST"])
def run_code():
    code = (request.get_json(force=True).get("code") or "").strip()
    if not code:
        return jsonify({"output": "(no code)"})
    return jsonify({"output": t_run_python(code)})


@app.route("/api/upload", methods=["POST"])
def upload():
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "No file received"}), 400
    name = f.filename or "file"
    ext = os.path.splitext(name)[1].lower()
    raw = f.read()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            if not text.strip():
                try:
                    text = ocr_pdf_bytes(raw)
                except Exception:
                    text = ""
        elif ext == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheets = []
            budget = 60_000
            for worksheet in workbook.worksheets:
                prefix = f"[Sheet: {worksheet.title}]\n"
                buffer = io.StringIO(newline="")
                writer = csv.writer(buffer, lineterminator="\n")
                for row in worksheet.iter_rows(values_only=True):
                    writer.writerow(
                        ["" if value is None else str(value) for value in row]
                    )
                    if len(prefix) + buffer.tell() >= budget:
                        break
                sheet_text = prefix + buffer.getvalue()
                sheets.append(sheet_text)
                budget -= len(sheet_text)
                if budget <= 0:
                    break
            workbook.close()
            text = "\n".join(sheets)
        elif ext == ".docx":
            from docx import Document

            document = Document(io.BytesIO(raw))
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
            for index, table in enumerate(document.tables, 1):
                buffer = io.StringIO(newline="")
                writer = csv.writer(buffer, lineterminator="\n")
                for row in table.rows:
                    writer.writerow(
                        [cell.text.replace("\n", " ").strip() for cell in row.cells]
                    )
                parts.append(f"[Table: Table {index}]\n{buffer.getvalue()}")
            text = "\n\n".join(parts)
        else:
            text = raw.decode("utf-8", errors="replace")
    except Exception as e:
        return jsonify({"error": f"Could not read file: {e}"}), 400
    limit = 60_000
    return jsonify({"name": name, "text": text[:limit], "truncated": len(text) > limit})


if __name__ == "__main__":
    app.run(debug=True, port=5000)
